from __future__ import annotations
import os, sys, re
from pathlib import Path
from textwrap import dedent
from dotenv import load_dotenv


_THIS = Path(__file__).resolve()
for parent in [_THIS.parent, *_THIS.parents]:
    if (parent / "helpers").is_dir():
        if str(parent) not in sys.path:
            sys.path.insert(0, str(parent))
        break
else:
    print("Could not locate 'helpers' directory on sys.path.", file=sys.stderr)
    sys.exit(1)


# Env + Config
load_dotenv()  

try:
    from helpers.config import load_config
except Exception as e:
    print(f"[config] Failed to load helpers.config: {e}", file=sys.stderr)
    sys.exit(1)

cfg = load_config()

# OpenAI SDK
try:
    from openai import OpenAI
except ImportError:
    print("Please install OpenAI SDK: pip install openai", file=sys.stderr)
    sys.exit(1)

# Optional .docx support
try:
    from docx import Document
    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False

# CONFIG (from config.yaml)
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
if not OPENAI_API_KEY:
    print("Missing OPENAI_API_KEY in .env", file=sys.stderr)
    sys.exit(3)

# Agent-specific model + temperature (with fallback to global)
MODEL = cfg.model_for("transcript_to_csv")
TEMPERATURE = cfg.temp_for("transcript_to_csv")
SYSTEM_PROMPT = dedent(cfg.prompt("transcript_to_csv", "system"))

# Helpers
def read_transcript(path: Path) -> str:
    """Reads .docx (including tables) or plain .txt into a single text blob."""
    if path.suffix.lower() != ".docx":
        return path.read_text(encoding="utf-8", errors="ignore")

    if not HAVE_DOCX:
        raise RuntimeError("Install python-docx first: pip install python-docx")

    doc = Document(str(path))
    parts = []

    # paragraphs
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)

    # tables
    for ti, table in enumerate(doc.tables, start=1):
        parts.append(f"\n=== TABLE {ti} ===")
        for ri, row in enumerate(table.rows, start=1):
            cells = []
            for ci, cell in enumerate(row.cells, start=1):
                cell_text = "\n".join(par.text.strip() for par in cell.paragraphs if par.text.strip())
                cell_text = re.sub(r"\s+", " ", cell_text).strip()
                cells.append(cell_text)
            parts.append(" | ".join(cells))
        parts.append(f"=== END TABLE {ti} ===\n")

    text = "\n".join(parts).strip()

    if len(text) < 200:
        text += ("\n\n[NOTE] The .docx appears to have very little extractable text. "
                 "If this is a scanned document or embedded image tables, convert it to .txt or a true .docx with real text.")
    return text

def run_agent(transcript_text: str, model: str | None = None) -> str:
    """Calls OpenAI Chat API using model/prompt from config.yaml (unless model override is provided)."""
    client = OpenAI(api_key=OPENAI_API_KEY)
    use_model = model or MODEL

    resp = client.chat.completions.create(
        model=use_model,
        temperature=TEMPERATURE,
        messages=[
            {"role": "system", "content": SYSTEM_PROMPT},
            {
                "role": "user",
                "content": (
                    "Transcript text below. Produce the CSV exactly per the spec.\n\n"
                    f"<<<TRANSCRIPT_START\n{transcript_text}\nTRANSCRIPT_END>>>"
                ),
            },
        ],
    )
    return (resp.choices[0].message.content or "").strip()

# CLI entry
def main():
    base_dir = _THIS.parents[1] 
    transcripts_dir = base_dir / "data" / "transcripts"
    out_dir = base_dir / "data" / "out"
    out_dir.mkdir(parents=True, exist_ok=True)

    transcript_files = sorted([p for p in transcripts_dir.glob("*") if p.suffix.lower() in {".docx", ".txt"}])
    if not transcript_files:
        print(f"No transcript files found in {transcripts_dir}")
        sys.exit(0)

    print(f"Found {len(transcript_files)} transcript(s):")
    for p in transcript_files:
        print(f" - {p.name}")

    for infile in transcript_files:
        outfile = out_dir / (infile.stem + ".csv")
        if outfile.exists():
            print(f"Skipping {infile.name} (already converted).")
            continue

        print(f"\nProcessing {infile.name} ...")
        transcript_text = read_transcript(infile)

        try:
            csv_text = run_agent(transcript_text) 
            outfile.write_text(csv_text, encoding="utf-8")
            print(f"✓ Converted → {outfile.name}")
        except Exception as e:
            print(f"[error] Failed on {infile.name}: {e}")

    print("\nAll transcripts processed!")

if __name__ == "__main__":
    main()